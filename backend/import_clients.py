#!/usr/bin/env python3
"""
Импортер клиентов из XLSX файла.

Структура файла клиентов:
- Лист: Лист_1
- Строк: 131 | Колонок: 23
- Колонки:
  A: Наименование
  B: Код
  C: Бизнес-регион
  D: Дата регистрации
  E: Клиент
  F: Комментарий
  G: Поставщик
  H: Публичное наименование
  I: Основной менеджер
  J: Прочие отношения
  K: Обслуживается торговыми представителями
  L: Прочая информация (пустая)
  M: Перевозчик
  N: Шаблон этикетки (пустая)
  O: Юр/Физлицо
  P: Пол (пустая)
  Q: Дата рождения (пустая)
  R: Вариант отправки электронного чека (пустая)
  S: Зона доставки (пустая)
  T: Вид цен (пустая)
  U: Индивидуальный вид цены (пустая)
  V: Водитель
  W: Ак флаг спец цена

Использование:
    python import_clients.py <путь_к_файлу.xlsx>
"""

import sys
import json
from pathlib import Path
from typing import List, Dict, Optional
from dataclasses import dataclass, asdict
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("❌ Установите openpyxl:")
    print("   pip install openpyxl")
    sys.exit(1)


@dataclass
class Client:
    """Клиент"""
    name: str  # Наименование
    code: Optional[str] = None  # Код
    business_region: Optional[str] = None  # Бизнес-регион
    registration_date: Optional[str] = None  # Дата регистрации
    client: Optional[str] = None  # Клиент
    comment: Optional[str] = None  # Комментарий
    supplier: Optional[str] = None  # Поставщик
    public_name: Optional[str] = None  # Публичное наименование
    main_manager: Optional[str] = None  # Основной менеджер
    other_relations: Optional[str] = None  # Прочие отношения
    serviced_by_sales_reps: Optional[str] = None  # Обслуживается торговыми представителями
    other_info: Optional[str] = None  # Прочая информация
    carrier: Optional[str] = None  # Перевозчик
    label_template: Optional[str] = None  # Шаблон этикетки
    legal_entity_type: Optional[str] = None  # Юр/Физлицо
    gender: Optional[str] = None  # Пол
    birth_date: Optional[str] = None  # Дата рождения
    receipt_delivery_method: Optional[str] = None  # Вариант отправки электронного чека
    delivery_zone: Optional[str] = None  # Зона доставки
    price_type: Optional[str] = None  # Вид цен
    individual_price_type: Optional[str] = None  # Индивидуальный вид цены
    driver: Optional[str] = None  # Водитель
    special_price_flag: Optional[str] = None  # Ак флаг спец цена
    
    # Метаданные
    row_number: int = 0  # Номер строки в файле
    import_status: str = "pending"  # pending, success, error
    error_message: Optional[str] = None
    
    def to_dict(self) -> dict:
        """Преобразовать в словарь"""
        return asdict(self)
    
    def validate(self) -> List[str]:
        """Валидация данных"""
        errors = []
        
        if not self.name or not self.name.strip():
            errors.append("Наименование обязательно")
        
        # Проверка типа Юр/Физлицо
        if self.legal_entity_type:
            valid_types = ['Юр', 'Физ', 'Юр.', 'Физ.', 'Юридическое', 'Физическое']
            if self.legal_entity_type not in valid_types:
                errors.append(f"Недопустимый тип Юр/Физлицо: {self.legal_entity_type}")
        
        return errors


def import_clients(file_path: str) -> Dict:
    """
    Импортирует клиентов из XLSX файла.
    
    Args:
        file_path: Путь к XLSX файлу
        
    Returns:
        Словарь с результатами импорта
    """
    print(f"\n{'='*70}")
    print(f"👥 ИМПОРТ КЛИЕНТОВ")
    print(f"{'='*70}")
    print(f"📁 Файл: {file_path}")
    
    path = Path(file_path)
    
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")
    
    if not path.suffix.lower() == '.xlsx':
        raise ValueError(f"Файл должен быть .xlsx, получено: {path.suffix}")
    
    # Загружаем workbook
    print(f"📖 Загружаем workbook...")
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    
    # Получаем первый лист
    if not wb.sheetnames:
        raise ValueError("Файл не содержит листов")
    
    sheet_name = wb.sheetnames[0]
    print(f"📄 Лист: {sheet_name}")
    ws = wb[sheet_name]
    
    # Получаем заголовки
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)
    
    print(f"📋 Найдено колонок: {len(headers)}")
    for idx, header in enumerate(headers, 1):
        print(f"   {idx}. {header}")
    
    # Импортируем строки
    clients: List[Client] = []
    row_number = 2  # Начинаем со второй строки (первая - заголовки)
    
    print(f"\n🔄 Импортируем данные...")
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):  # Пропускаем пустые строки
            continue
        
        # Создаём клиента
        client = Client(
            name=str(row[0] or '').strip(),  # A: Наименование
            code=str(row[1] or '').strip() if len(row) > 1 else None,  # B: Код
            business_region=str(row[2] or '').strip() if len(row) > 2 else None,  # C: Бизнес-регион
            registration_date=str(row[3] or '').strip() if len(row) > 3 else None,  # D: Дата регистрации
            client=str(row[4] or '').strip() if len(row) > 4 else None,  # E: Клиент
            comment=str(row[5] or '').strip() if len(row) > 5 and row[5] else None,  # F: Комментарий
            supplier=str(row[6] or '').strip() if len(row) > 6 else None,  # G: Поставщик
            public_name=str(row[7] or '').strip() if len(row) > 7 else None,  # H: Публичное наименование
            main_manager=str(row[8] or '').strip() if len(row) > 8 else None,  # I: Основной менеджер
            other_relations=str(row[9] or '').strip() if len(row) > 9 else None,  # J: Прочие отношения
            serviced_by_sales_reps=str(row[10] or '').strip() if len(row) > 10 else None,  # K: Обслуживается торговыми представителями
            other_info=str(row[11] or '').strip() if len(row) > 11 and row[11] else None,  # L: Прочая информация
            carrier=str(row[12] or '').strip() if len(row) > 12 else None,  # M: Перевозчик
            label_template=str(row[13] or '').strip() if len(row) > 13 and row[13] else None,  # N: Шаблон этикетки
            legal_entity_type=str(row[14] or '').strip() if len(row) > 14 else None,  # O: Юр/Физлицо
            gender=str(row[15] or '').strip() if len(row) > 15 and row[15] else None,  # P: Пол
            birth_date=str(row[16] or '').strip() if len(row) > 16 and row[16] else None,  # Q: Дата рождения
            receipt_delivery_method=str(row[17] or '').strip() if len(row) > 17 and row[17] else None,  # R: Вариант отправки электронного чека
            delivery_zone=str(row[18] or '').strip() if len(row) > 18 and row[18] else None,  # S: Зона доставки
            price_type=str(row[19] or '').strip() if len(row) > 19 and row[19] else None,  # T: Вид цен
            individual_price_type=str(row[20] or '').strip() if len(row) > 20 and row[20] else None,  # U: Индивидуальный вид цены
            driver=str(row[21] or '').strip() if len(row) > 21 and row[21] else None,  # V: Водитель
            special_price_flag=str(row[22] or '').strip() if len(row) > 22 else None,  # W: Ак флаг спец цена
            row_number=row_number,
        )
        
        clients.append(client)
        row_number += 1
    
    wb.close()
    
    print(f"✅ Прочитано строк: {len(clients)}")
    
    # Валидация
    print(f"\n🔍 Валидация данных...")
    valid_clients = []
    invalid_clients = []
    
    for client in clients:
        errors = client.validate()
        
        if errors:
            client.import_status = "error"
            client.error_message = "; ".join(errors)
            invalid_clients.append(client)
            print(f"   ❌ Строка {client.row_number}: {client.error_message}")
        else:
            client.import_status = "success"
            valid_clients.append(client)
    
    print(f"✅ Валидных: {len(valid_clients)}")
    print(f"❌ Ошибок: {len(invalid_clients)}")
    
    # Статистика
    result = {
        'file': str(path),
        'file_size_mb': round(path.stat().st_size / 1024 / 1024, 2),
        'sheet_name': sheet_name,
        'total_rows': len(clients),
        'valid_rows': len(valid_clients),
        'invalid_rows': len(invalid_clients),
        'clients': [client.to_dict() for client in clients],
        'import_time': datetime.now().isoformat(),
    }
    
    # Сохраняем результаты
    output_file = path.with_suffix('.import.json')
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    
    print(f"\n💾 Результаты сохранены: {output_file}")
    
    # Статистика по полям
    print(f"\n📊 Статистика по полям:")
    print(f"   Наименование: {sum(1 for c in clients if c.name)} заполнено")
    print(f"   Код: {sum(1 for c in clients if c.code)} заполнено")
    print(f"   Бизнес-регион: {sum(1 for c in clients if c.business_region)} заполнено")
    print(f"   Дата регистрации: {sum(1 for c in clients if c.registration_date)} заполнено")
    print(f"   Основной менеджер: {sum(1 for c in clients if c.main_manager)} заполнено")
    print(f"   Перевозчик: {sum(1 for c in clients if c.carrier)} заполнено")
    print(f"   Водитель: {sum(1 for c in clients if c.driver)} заполнено")
    print(f"   Юр/Физлицо: {sum(1 for c in clients if c.legal_entity_type)} заполнено")
    
    print(f"\n{'='*70}")
    print(f"✅ ИМПОРТ ЗАВЕРШЁН")
    print(f"{'='*70}\n")
    
    return result


def main():
    """Главная функция"""
    if len(sys.argv) < 2:
        print("Использование: python import_clients.py <путь_к_файлу.xlsx>")
        print()
        print("Пример:")
        print("  python import_clients.py Клиенты.xlsx")
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    try:
        result = import_clients(file_path)
        
        print(f"👥 Импортировано: {result['valid_rows']} из {result['total_rows']} клиентов")
        
        if result['invalid_rows'] > 0:
            print(f"⚠️  Ошибок валидации: {result['invalid_rows']}")
            print(f"   Подробности в файле: {Path(file_path).with_suffix('.import.json')}")
        
    except Exception as e:
        print(f"\n❌ Ошибка: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
