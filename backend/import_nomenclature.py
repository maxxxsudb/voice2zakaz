#!/usr/bin/env python3
"""
Импортер номенклатуры из XLSX файла.

Структура файла "ЦыганковНоменклатура.xlsx":
- Лист: Лист_1
- Строк: 318 | Колонок: 11
- Колонки:
  A: Наименование
  B: Артикул
  C: Единица измерения веса
  D: Вес (знаменатель)
  E: Вес
  F: Вес (числитель)
  G: Вид номенклатуры
  H: Единица для отчетов
  I: Единица хранения
  J: GTIN (пустая)
  K: Код

Использование:
    python import_nomenclature.py <путь_к_файлу.xlsx>
"""

import sys
import json
from pathlib import Path
from typing import List, Dict, Optional
from dataclasses import dataclass, asdict
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("❌ Установите openpyxl:")
    print("   pip install openpyxl")
    sys.exit(1)


@dataclass
class NomenclatureItem:
    """Элемент номенклатуры"""
    name: str  # Наименование
    article: Optional[str] = None  # Артикул
    weight_unit: Optional[str] = None  # Единица измерения веса
    weight_denominator: Optional[float] = None  # Вес (знаменатель)
    weight: Optional[str] = None  # Вес
    weight_numerator: Optional[float] = None  # Вес (числитель)
    nomenclature_type: Optional[str] = None  # Вид номенклатуры
    report_unit: Optional[str] = None  # Единица для отчетов
    storage_unit: Optional[str] = None  # Единица хранения
    gtin: Optional[str] = None  # GTIN
    code: Optional[str] = None  # Код
    
    # Метаданные
    row_number: int = 0  # Номер строки в файле
    import_status: str = "pending"  # pending, success, error
    error_message: Optional[str] = None
    
    def to_dict(self) -> dict:
        """Преобразовать в словарь"""
        return asdict(self)
    
    def validate(self) -> List[str]:
        """Валидация данных"""
        errors = []
        
        if not self.name or not self.name.strip():
            errors.append("Наименование обязательно")
        
        if self.weight_denominator is not None and self.weight_denominator <= 0:
            errors.append("Вес (знаменатель) должен быть > 0")
        
        if self.weight_numerator is not None and self.weight_numerator <= 0:
            errors.append("Вес (числитель) должен быть > 0")
        
        return errors


def import_nomenclature(file_path: str) -> Dict:
    """
    Импортирует номенклатуру из XLSX файла.
    
    Args:
        file_path: Путь к XLSX файлу
        
    Returns:
        Словарь с результатами импорта
    """
    print(f"\n{'='*70}")
    print(f"📦 ИМПОРТ НОМЕНКЛАТУРЫ")
    print(f"{'='*70}")
    print(f"📁 Файл: {file_path}")
    
    path = Path(file_path)
    
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")
    
    if not path.suffix.lower() == '.xlsx':
        raise ValueError(f"Файл должен быть .xlsx, получено: {path.suffix}")
    
    # Загружаем workbook
    print(f"📖 Загружаем workbook...")
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    
    # Получаем первый лист
    if not wb.sheetnames:
        raise ValueError("Файл не содержит листов")
    
    sheet_name = wb.sheetnames[0]
    print(f"📄 Лист: {sheet_name}")
    ws = wb[sheet_name]
    
    # Получаем заголовки
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)
    
    print(f"📋 Найдено колонок: {len(headers)}")
    for idx, header in enumerate(headers, 1):
        print(f"   {idx}. {header}")
    
    # Импортируем строки
    items: List[NomenclatureItem] = []
    row_number = 2  # Начинаем со второй строки (первая - заголовки)
    
    print(f"\n🔄 Импортируем данные...")
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):  # Пропускаем пустые строки
            continue
        
        # Создаём элемент номенклатуры
        item = NomenclatureItem(
            name=str(row[0] or '').strip(),  # A: Наименование
            article=str(row[1] or '').strip() if len(row) > 1 else None,  # B: Артикул
            weight_unit=str(row[2] or '').strip() if len(row) > 2 else None,  # C: Единица измерения веса
            weight_denominator=float(row[3]) if len(row) > 3 and row[3] is not None else None,  # D: Вес (знаменатель)
            weight=str(row[4] or '').strip() if len(row) > 4 else None,  # E: Вес
            weight_numerator=float(row[5]) if len(row) > 5 and row[5] is not None else None,  # F: Вес (числитель)
            nomenclature_type=str(row[6] or '').strip() if len(row) > 6 else None,  # G: Вид номенклатуры
            report_unit=str(row[7] or '').strip() if len(row) > 7 else None,  # H: Единица для отчетов
            storage_unit=str(row[8] or '').strip() if len(row) > 8 else None,  # I: Единица хранения
            gtin=str(row[9] or '').strip() if len(row) > 9 and row[9] else None,  # J: GTIN
            code=str(row[10] or '').strip() if len(row) > 10 else None,  # K: Код
            row_number=row_number,
        )
        
        items.append(item)
        row_number += 1
    
    wb.close()
    
    print(f"✅ Прочитано строк: {len(items)}")
    
    # Валидация
    print(f"\n🔍 Валидация данных...")
    valid_items = []
    invalid_items = []
    
    for item in items:
        errors = item.validate()
        
        if errors:
            item.import_status = "error"
            item.error_message = "; ".join(errors)
            invalid_items.append(item)
            print(f"   ❌ Строка {item.row_number}: {item.error_message}")
        else:
            item.import_status = "success"
            valid_items.append(item)
    
    print(f"✅ Валидных: {len(valid_items)}")
    print(f"❌ Ошибок: {len(invalid_items)}")
    
    # Статистика
    result = {
        'file': str(path),
        'file_size_mb': round(path.stat().st_size / 1024 / 1024, 2),
        'sheet_name': sheet_name,
        'total_rows': len(items),
        'valid_rows': len(valid_items),
        'invalid_rows': len(invalid_items),
        'items': [item.to_dict() for item in items],
        'import_time': datetime.now().isoformat(),
    }
    
    # Сохраняем результаты
    output_file = path.with_suffix('.import.json')
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    
    print(f"\n💾 Результаты сохранены: {output_file}")
    
    # Статистика по полям
    print(f"\n📊 Статистика по полям:")
    print(f"   Наименование: {sum(1 for i in items if i.name)} заполнено")
    print(f"   Артикул: {sum(1 for i in items if i.article)} заполнено")
    print(f"   Вес (числитель): {sum(1 for i in items if i.weight_numerator is not None)} заполнено")
    print(f"   Вес (знаменатель): {sum(1 for i in items if i.weight_denominator is not None)} заполнено")
    print(f"   Вид номенклатуры: {sum(1 for i in items if i.nomenclature_type)} заполнено")
    print(f"   Код: {sum(1 for i in items if i.code)} заполнено")
    
    print(f"\n{'='*70}")
    print(f"✅ ИМПОРТ ЗАВЕРШЁН")
    print(f"{'='*70}\n")
    
    return result


def main():
    """Главная функция"""
    if len(sys.argv) < 2:
        print("Использование: python import_nomenclature.py <путь_к_файлу.xlsx>")
        print()
        print("Пример:")
        print("  python import_nomenclature.py ЦыганковНоменклатура.xlsx")
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    try:
        result = import_nomenclature(file_path)
        
        print(f"📦 Импортировано: {result['valid_rows']} из {result['total_rows']} записей")
        
        if result['invalid_rows'] > 0:
            print(f"⚠️  Ошибок валидации: {result['invalid_rows']}")
            print(f"   Подробности в файле: {Path(file_path).with_suffix('.import.json')}")
        
    except Exception as e:
        print(f"\n❌ Ошибка: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
