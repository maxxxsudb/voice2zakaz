#!/usr/bin/env python3
"""
Анализатор XLSX файлов.
Читает Excel файл и показывает структуру данных:
- Названия колонок
- Типы данных
- Примеры значений
- Количество строк
- Статистику по каждой колонке

Использование:
    python analyze_xlsx.py <путь_к_файлу.xlsx>

Пример:
    python analyze_xlsx.py nomenclature.xlsx
    python analyze_xlsx.py clients.xlsx
"""

import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    print("❌ Установите openpyxl:")
    print("   pip install openpyxl")
    sys.exit(1)


def analyze_xlsx(file_path: str) -> dict:
    """
    Анализирует XLSX файл и возвращает информацию о структуре.
    
    Args:
        file_path: Путь к XLSX файлу
        
    Returns:
        Словарь с информацией о файле
    """
    print(f"\n🔍 [XLSX MODULE] Начинаем анализ файла: {file_path}")
    
    path = Path(file_path)
    
    if not path.exists():
        print(f"❌ [XLSX MODULE] Файл не найден: {file_path}")
        raise FileNotFoundError(f"Файл не найден: {file_path}")
    
    if not path.suffix.lower() == '.xlsx':
        print(f"❌ [XLSX MODULE] Неправильное расширение: {path.suffix}")
        raise ValueError(f"Файл должен быть .xlsx, получено: {path.suffix}")
    
    print(f"✅ [XLSX MODULE] Файл существует, размер: {path.stat().st_size} байт")
    
    # Загружаем workbook
    print("📖 [XLSX MODULE] Загружаем workbook...")
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    print(f"✅ [XLSX MODULE] Workbook загружен, листов: {len(wb.sheetnames)}")
    
    result = {
        'file': str(path),
        'file_size_mb': round(path.stat().st_size / 1024 / 1024, 2),
        'sheets': [],
    }
    
    # Анализируем каждый лист
    for sheet_idx, sheet_name in enumerate(wb.sheetnames, 1):
        print(f"\n📄 [XLSX MODULE] Анализируем лист {sheet_idx}/{len(wb.sheetnames)}: {sheet_name}")
        ws = wb[sheet_name]
        
        sheet_info = {
            'name': sheet_name,
            'max_row': ws.max_row,
            'max_column': ws.max_column,
            'columns': [],
            'sample_rows': [],
        }
        
        print(f"   Строк: {ws.max_row}, Колонок: {ws.max_column}")
        
        # Получаем заголовки (первая строка)
        print("   🔍 Читаем заголовки...")
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        print(f"   ✅ Заголовков: {len(headers)}")
        
        # Анализируем каждую колонку
        print("   🔍 Анализируем колонки...")
        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            
            # Собираем значения из колонки (первые 100 строк для анализа)
            values = []
            for row_idx in range(2, min(ws.max_row + 1, 102)):  # Пропускаем заголовок
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    values.append(cell.value)
            
            # Определяем тип данных
            types = set()
            for v in values:
                if isinstance(v, (int, float)):
                    types.add('number')
                elif isinstance(v, bool):
                    types.add('boolean')
                else:
                    types.add('string')
            
            # Статистика для числовых колонок
            stats = {}
            if types == {'number'} and values:
                numeric_values = [v for v in values if isinstance(v, (int, float))]
                if numeric_values:
                    stats = {
                        'min': min(numeric_values),
                        'max': max(numeric_values),
                        'avg': sum(numeric_values) / len(numeric_values),
                    }
            
            # Преобразуем значения в строки для JSON
            sample_values_str = []
            for v in values[:5]:
                if v is None:
                    sample_values_str.append(None)
                elif isinstance(v, (int, float, bool)):
                    sample_values_str.append(v)
                else:
                    sample_values_str.append(str(v))
            
            col_info = {
                'letter': col_letter,
                'header': str(header) if header else None,
                'type': ', '.join(sorted(types)) if types else 'empty',
                'non_empty_count': len(values),
                'sample_values': sample_values_str,
                'stats': stats,
            }
            
            sheet_info['columns'].append(col_info)
            print(f"      {col_letter}: {header or '(без заголовка)'} [{col_info['type']}] ({len(values)} значений)")
        
        # Добавляем примеры строк (первые 3)
        print("   🔍 Собираем примеры строк...")
        for row_idx in range(2, min(ws.max_row + 1, 5)):
            row_data = {}
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = cell.value
                # Преобразуем в строку для JSON
                if value is None:
                    row_data[str(header) if header else f'Column_{col_idx}'] = None
                elif isinstance(value, (int, float, bool)):
                    row_data[str(header) if header else f'Column_{col_idx}'] = value
                else:
                    row_data[str(header) if header else f'Column_{col_idx}'] = str(value)
            sheet_info['sample_rows'].append(row_data)
        
        print(f"   ✅ Примеров строк: {len(sheet_info['sample_rows'])}")
        
        result['sheets'].append(sheet_info)
    
    wb.close()
    print(f"\n✅ [XLSX MODULE] Анализ завершён, листов обработано: {len(result['sheets'])}")
    return result


def print_analysis(result: dict):
    """Выводит результаты анализа в читаемом формате"""
    print("=" * 70)
    print(f"📊 АНАЛИЗ ФАЙЛА: {result['file']}")
    print("=" * 70)
    print(f"Размер файла: {result['file_size_mb']} МБ")
    print(f"Количество листов: {len(result['sheets'])}")
    print()
    
    for sheet in result['sheets']:
        print("-" * 70)
        print(f"📄 ЛИСТ: {sheet['name']}")
        print(f"   Строк: {sheet['max_row']} | Колонок: {sheet['max_column']}")
        print()
        
        # Таблица колонок
        print("📋 КОЛОНКИ:")
        print(f"{'Буква':<6} {'Заголовок':<30} {'Тип':<15} {'Заполнено':<10}")
        print("-" * 70)
        
        for col in sheet['columns']:
            header = str(col['header'] or '(пусто)')[:28]
            print(f"{col['letter']:<6} {header:<30} {col['type']:<15} {col['non_empty_count']:<10}")
            
            # Статистика для числовых колонок
            if col['stats']:
                stats = col['stats']
                print(f"       └─ Min: {stats['min']}, Max: {stats['max']}, Avg: {stats['avg']:.2f}")
        
        print()
        
        # Примеры значений
        print("🔍 ПРИМЕРЫ ЗНАЧЕНИЙ (первые 5 из каждой колонки):")
        for col in sheet['columns']:
            if col['sample_values']:
                samples = [str(v)[:40] for v in col['sample_values']]
                print(f"   {col['header'] or col['letter']}: {', '.join(samples)}")
        
        print()
        
        # Примеры строк
        if sheet['sample_rows']:
            print("📝 ПРИМЕРЫ СТРОК (первые 3):")
            for idx, row in enumerate(sheet['sample_rows'], start=1):
                print(f"   Строка {idx}:")
                for key, value in row.items():
                    value_str = str(value)[:50] if value is not None else '(пусто)'
                    print(f"      {key}: {value_str}")
            print()


def main():
    """Главная функция"""
    if len(sys.argv) < 2:
        print("Использование: python analyze_xlsx.py <путь_к_файлу.xlsx>")
        print()
        print("Пример:")
        print("  python analyze_xlsx.py nomenclature.xlsx")
        print("  python analyze_xlsx.py clients.xlsx")
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    try:
        print(f"\n⏳ Анализирую файл: {file_path}\n")
        result = analyze_xlsx(file_path)
        print_analysis(result)
        
        print("=" * 70)
        print("✅ Анализ завершён!")
        print("=" * 70)
        print()
        print("💡 На основе этой информации можно написать импортер.")
        print("   Скопируйте вывод и передайте его для генерации кода импорта.")
        
    except FileNotFoundError as e:
        print(f"❌ {e}")
        sys.exit(1)
    except Exception as e:
        print(f"❌ Ошибка: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
