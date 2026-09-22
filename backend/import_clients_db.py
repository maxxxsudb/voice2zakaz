#!/usr/bin/env python3
"""
Импортер клиентов с сохранением в PostgreSQL.

Использование:
    python import_clients_db.py <employee_id> <путь_к_файлу.xlsx>

Пример:
    python import_clients_db.py ivanov Клиенты.xlsx
"""

import sys
import uuid
from pathlib import Path
from typing import List
from repositories import EmployeeRepository, ClientRepository, VoiceDictionaryRepository

try:
    import openpyxl
except ImportError:
    print("❌ Установите openpyxl:")
    print("   pip install openpyxl")
    sys.exit(1)


def import_clients(employee_id: str, file_path: str) -> dict:
    """
    Импортирует клиентов из XLSX файла и сохраняет в PostgreSQL.
    
    Args:
        employee_id: ID сотрудника
        file_path: Путь к XLSX файлу
        
    Returns:
        Словарь с результатами импорта
    """
    print(f"\n{'='*70}")
    print(f"👥 ИМПОРТ КЛИЕНТОВ ДЛЯ СОТРУДНИКА: {employee_id}")
    print(f"{'='*70}")
    print(f"📁 Файл: {file_path}")
    
    path = Path(file_path)
    
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")
    
    if not path.suffix.lower() == '.xlsx':
        raise ValueError(f"Файл должен быть .xlsx, получено: {path.suffix}")
    
    # Получаем или создаем сотрудника
    employee = EmployeeRepository.get_by_id(employee_id)
    if not employee:
        print(f"⚠️  Сотрудник {employee_id} не найден, создаем нового...")
        employee = EmployeeRepository.create(
            employee_id=employee_id,
            name=employee_id
        )
    
    print(f"✅ Сотрудник: {employee.name}")
    
    # Загружаем workbook
    print(f"\n📖 Загружаем workbook...")
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    
    # Получаем первый лист
    if not wb.sheetnames:
        raise ValueError("Файл не содержит листов")
    
    sheet_name = wb.sheetnames[0]
    print(f"📄 Лист: {sheet_name}")
    ws = wb[sheet_name]
    
    # Получаем заголовки
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)
    
    print(f"📋 Найдено колонок: {len(headers)}")
    
    # Импортируем строки
    clients_data = []
    row_number = 2
    
    print(f"\n🔄 Импортируем данные...")
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        
        # Создаем уникальный ID
        client_id = f"c{uuid.uuid4().hex[:8]}"
        
        # Создаем данные для клиента
        client_data = {
            'id': client_id,
            'employee_id': employee_id,
            'name': str(row[0] or '').strip(),
            'code': str(row[1] or '').strip() if len(row) > 1 and row[1] else None,
            'business_region': str(row[2] or '').strip() if len(row) > 2 and row[2] else None,
            'registration_date': str(row[3] or '').strip() if len(row) > 3 and row[3] else None,
            'client_type': str(row[4] or '').strip() if len(row) > 4 and row[4] else None,
            'comment': str(row[5] or '').strip() if len(row) > 5 and row[5] else None,
            'supplier': str(row[6] or '').strip() if len(row) > 6 and row[6] else None,
            'public_name': str(row[7] or '').strip() if len(row) > 7 and row[7] else None,
            'main_manager': str(row[8] or '').strip() if len(row) > 8 and row[8] else None,
            'other_relations': str(row[9] or '').strip() if len(row) > 9 and row[9] else None,
            'serviced_by_sales_reps': str(row[10] or '').strip() if len(row) > 10 and row[10] else None,
            'carrier': str(row[12] or '').strip() if len(row) > 12 and row[12] else None,
            'legal_entity_type': str(row[14] or '').strip() if len(row) > 14 and row[14] else None,
            'driver': str(row[21] or '').strip() if len(row) > 21 and row[21] else None,
            'special_price_flag': str(row[22] or '').strip() if len(row) > 22 and row[22] else None,
            'row_number': row_number,
            'import_status': 'pending',
        }
        
        clients_data.append(client_data)
        row_number += 1
    
    wb.close()
    
    print(f"✅ Прочитано строк: {len(clients_data)}")
    
    # Валидация
    print(f"\n🔍 Валидация данных...")
    valid_clients = []
    invalid_clients = []
    
    for client_data in clients_data:
        errors = []
        
        if not client_data['name'] or not client_data['name'].strip():
            errors.append("Наименование обязательно")
        
        if errors:
            client_data['import_status'] = 'error'
            client_data['error_message'] = "; ".join(errors)
            invalid_clients.append(client_data)
            print(f"   ❌ Строка {client_data['row_number']}: {client_data['error_message']}")
        else:
            client_data['import_status'] = 'success'
            valid_clients.append(client_data)
    
    print(f"✅ Валидных: {len(valid_clients)}")
    print(f"❌ Ошибок: {len(invalid_clients)}")
    
    # Сохраняем в БД
    print(f"\n💾 Сохраняем в базу данных...")
    
    # Сохраняем все клиенты (включая с ошибками)
    all_clients = ClientRepository.bulk_create(clients_data)
    
    # Создаем записи в словаре для валидных клиентов
    print(f"\n📖 Создаем словарь для распознавания...")
    dictionary_count = 0
    
    for client_data in valid_clients:
        # Создаем запись в словаре
        entry = VoiceDictionaryRepository.create(
            employee_id=employee_id,
            original=client_data['name'],
            category='client',
            item_id=client_data['id']
        )
        
        # Добавляем код и публичное название как варианты
        if client_data['code']:
            VoiceDictionaryRepository.add_variant(entry.id, client_data['code'])
        if client_data['public_name'] and client_data['public_name'] != client_data['name']:
            VoiceDictionaryRepository.add_variant(entry.id, client_data['public_name'])
        
        dictionary_count += 1
    
    print(f"✅ Создано записей в словаре: {dictionary_count}")
    
    # Статистика
    result = {
        'employee_id': employee_id,
        'employee_name': employee.name,
        'file': str(path),
        'file_size_mb': round(path.stat().st_size / 1024 / 1024, 2),
        'sheet_name': sheet_name,
        'total_rows': len(clients_data),
        'valid_rows': len(valid_clients),
        'invalid_rows': len(invalid_clients),
        'dictionary_entries': dictionary_count,
    }
    
    # Статистика по полям
    print(f"\n📊 Статистика по полям:")
    print(f"   Наименование: {sum(1 for c in clients_data if c['name'])} заполнено")
    print(f"   Код: {sum(1 for c in clients_data if c['code'])} заполнено")
    print(f"   Бизнес-регион: {sum(1 for c in clients_data if c['business_region'])} заполнено")
    print(f"   Основной менеджер: {sum(1 for c in clients_data if c['main_manager'])} заполнено")
    print(f"   Перевозчик: {sum(1 for c in clients_data if c['carrier'])} заполнено")
    print(f"   Водитель: {sum(1 for c in clients_data if c['driver'])} заполнено")
    
    print(f"\n{'='*70}")
    print(f"✅ ИМПОРТ ЗАВЕРШЁН")
    print(f"{'='*70}\n")
    
    return result


def main():
    """Главная функция"""
    if len(sys.argv) < 3:
        print("Использование: python import_clients_db.py <employee_id> <путь_к_файлу.xlsx>")
        print()
        print("Пример:")
        print("  python import_clients_db.py ivanov Клиенты.xlsx")
        sys.exit(1)
    
    employee_id = sys.argv[1]
    file_path = sys.argv[2]
    
    try:
        result = import_clients(employee_id, file_path)
        
        print(f"👥 Импортировано: {result['valid_rows']} из {result['total_rows']} клиентов")
        print(f"👤 Сотрудник: {result['employee_name']}")
        print(f"📖 Записей в словаре: {result['dictionary_entries']}")
        
        if result['invalid_rows'] > 0:
            print(f"⚠️  Ошибок валидации: {result['invalid_rows']}")
        
    except Exception as e:
        print(f"\n❌ Ошибка: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
