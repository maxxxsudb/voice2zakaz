#!/usr/bin/env python3
"""
Импортер номенклатуры с сохранением в PostgreSQL.

Использование:
    python import_nomenclature_db.py <employee_id> <путь_к_файлу.xlsx>

Пример:
    python import_nomenclature_db.py ivanov ЦыганковНоменклатура.xlsx
"""

import sys
import uuid
from pathlib import Path
from typing import List
from repositories import EmployeeRepository, NomenclatureRepository, VoiceDictionaryRepository

try:
    import openpyxl
except ImportError:
    print("❌ Установите openpyxl:")
    print("   pip install openpyxl")
    sys.exit(1)


def import_nomenclature(employee_id: str, file_path: str) -> dict:
    """
    Импортирует номенклатуру из XLSX файла и сохраняет в PostgreSQL.
    
    Args:
        employee_id: ID сотрудника
        file_path: Путь к XLSX файлу
        
    Returns:
        Словарь с результатами импорта
    """
    print(f"\n{'='*70}")
    print(f"📦 ИМПОРТ НОМЕНКЛАТУРЫ ДЛЯ СОТРУДНИКА: {employee_id}")
    print(f"{'='*70}")
    print(f"📁 Файл: {file_path}")
    
    path = Path(file_path)
    
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")
    
    if not path.suffix.lower() == '.xlsx':
        raise ValueError(f"Файл должен быть .xlsx, получено: {path.suffix}")
    
    # Получаем или создаем сотрудника
    employee = EmployeeRepository.get_by_id(employee_id)
    if not employee:
        print(f"⚠️  Сотрудник {employee_id} не найден, создаем нового...")
        employee = EmployeeRepository.create(
            employee_id=employee_id,
            name=employee_id
        )
    
    print(f"✅ Сотрудник: {employee.name}")
    
    # Загружаем workbook
    print(f"\n📖 Загружаем workbook...")
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    
    # Получаем первый лист
    if not wb.sheetnames:
        raise ValueError("Файл не содержит листов")
    
    sheet_name = wb.sheetnames[0]
    print(f"📄 Лист: {sheet_name}")
    ws = wb[sheet_name]
    
    # Получаем заголовки
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)
    
    print(f"📋 Найдено колонок: {len(headers)}")
    
    # Импортируем строки
    items_data = []
    row_number = 2
    
    print(f"\n🔄 Импортируем данные...")
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        
        # Создаем уникальный ID
        item_id = f"n{uuid.uuid4().hex[:8]}"
        
        # Создаем данные для номенклатуры
        item_data = {
            'id': item_id,
            'employee_id': employee_id,
            'name': str(row[0] or '').strip(),
            'article': str(row[1] or '').strip() if len(row) > 1 and row[1] else None,
            'weight_unit': str(row[2] or '').strip() if len(row) > 2 and row[2] else None,
            'weight_denominator': float(row[3]) if len(row) > 3 and row[3] is not None else None,
            'weight': str(row[4] or '').strip() if len(row) > 4 and row[4] else None,
            'weight_numerator': float(row[5]) if len(row) > 5 and row[5] is not None else None,
            'nomenclature_type': str(row[6] or '').strip() if len(row) > 6 and row[6] else None,
            'report_unit': str(row[7] or '').strip() if len(row) > 7 and row[7] else None,
            'storage_unit': str(row[8] or '').strip() if len(row) > 8 and row[8] else None,
            'gtin': str(row[9] or '').strip() if len(row) > 9 and row[9] else None,
            'code': str(row[10] or '').strip() if len(row) > 10 and row[10] else None,
            'row_number': row_number,
            'import_status': 'pending',
        }
        
        items_data.append(item_data)
        row_number += 1
    
    wb.close()
    
    print(f"✅ Прочитано строк: {len(items_data)}")
    
    # Валидация
    print(f"\n🔍 Валидация данных...")
    valid_items = []
    invalid_items = []
    
    for item_data in items_data:
        errors = []
        
        if not item_data['name'] or not item_data['name'].strip():
            errors.append("Наименование обязательно")
        
        if item_data['weight_denominator'] is not None and item_data['weight_denominator'] <= 0:
            errors.append("Вес (знаменатель) должен быть > 0")
        
        if item_data['weight_numerator'] is not None and item_data['weight_numerator'] <= 0:
            errors.append("Вес (числитель) должен быть > 0")
        
        if errors:
            item_data['import_status'] = 'error'
            item_data['error_message'] = "; ".join(errors)
            invalid_items.append(item_data)
            print(f"   ❌ Строка {item_data['row_number']}: {item_data['error_message']}")
        else:
            item_data['import_status'] = 'success'
            valid_items.append(item_data)
    
    print(f"✅ Валидных: {len(valid_items)}")
    print(f"❌ Ошибок: {len(invalid_items)}")
    
    # Сохраняем в БД
    print(f"\n💾 Сохраняем в базу данных...")
    
    # Сохраняем все элементы (включая с ошибками)
    all_items = NomenclatureRepository.bulk_create(items_data)
    
    # Создаем записи в словаре для валидных элементов
    print(f"\n📖 Создаем словарь для распознавания...")
    dictionary_count = 0
    
    for item_data in valid_items:
        # Создаем запись в словаре
        entry = VoiceDictionaryRepository.create(
            employee_id=employee_id,
            original=item_data['name'],
            category='nomenclature',
            item_id=item_data['id']
        )
        
        # Добавляем артикул и код как варианты
        if item_data['article']:
            VoiceDictionaryRepository.add_variant(entry.id, item_data['article'])
        if item_data['code']:
            VoiceDictionaryRepository.add_variant(entry.id, item_data['code'])
        
        dictionary_count += 1
    
    print(f"✅ Создано записей в словаре: {dictionary_count}")
    
    # Статистика
    result = {
        'employee_id': employee_id,
        'employee_name': employee.name,
        'file': str(path),
        'file_size_mb': round(path.stat().st_size / 1024 / 1024, 2),
        'sheet_name': sheet_name,
        'total_rows': len(items_data),
        'valid_rows': len(valid_items),
        'invalid_rows': len(invalid_items),
        'dictionary_entries': dictionary_count,
    }
    
    # Статистика по полям
    print(f"\n📊 Статистика по полям:")
    print(f"   Наименование: {sum(1 for i in items_data if i['name'])} заполнено")
    print(f"   Артикул: {sum(1 for i in items_data if i['article'])} заполнено")
    print(f"   Код: {sum(1 for i in items_data if i['code'])} заполнено")
    print(f"   Вес (числитель): {sum(1 for i in items_data if i['weight_numerator'] is not None)} заполнено")
    print(f"   Вид номенклатуры: {sum(1 for i in items_data if i['nomenclature_type'])} заполнено")
    
    print(f"\n{'='*70}")
    print(f"✅ ИМПОРТ ЗАВЕРШЁН")
    print(f"{'='*70}\n")
    
    return result


def main():
    """Главная функция"""
    if len(sys.argv) < 3:
        print("Использование: python import_nomenclature_db.py <employee_id> <путь_к_файлу.xlsx>")
        print()
        print("Пример:")
        print("  python import_nomenclature_db.py ivanov ЦыганковНоменклатура.xlsx")
        sys.exit(1)
    
    employee_id = sys.argv[1]
    file_path = sys.argv[2]
    
    try:
        result = import_nomenclature(employee_id, file_path)
        
        print(f"📦 Импортировано: {result['valid_rows']} из {result['total_rows']} записей")
        print(f"👤 Сотрудник: {result['employee_name']}")
        print(f"📖 Записей в словаре: {result['dictionary_entries']}")
        
        if result['invalid_rows'] > 0:
            print(f"⚠️  Ошибок валидации: {result['invalid_rows']}")
        
    except Exception as e:
        print(f"\n❌ Ошибка: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
